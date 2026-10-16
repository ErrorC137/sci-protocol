import io

from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.forms.mapper import PLACEHOLDER, PLACEHOLDER_TH, PLACEHOLDERS, has_placeholder


def _create(client, persona, grant_id, template_id=None):
    body = {"grant_id": grant_id, "profile": persona["profile"], "applicant": persona["applicant"]}
    if template_id:
        body["template_id"] = template_id
    r = client.post("/applications", json=body)
    assert r.status_code == 200, r.text
    return r.json()


def _resolve_everything(client, app_id, fields):
    """What the researcher does on the review screen."""
    edits = []
    for f in fields:
        if has_placeholder(f["value"]):
            head = f["value"]
            for ph in PLACEHOLDERS:
                head = head.split(ph)[0]
            edits.append({"field_id": f["field_id"], "value": head + "Researcher-written text."})
        elif f["status"] == "amber":
            edits.append({"field_id": f["field_id"], "confirmed": True})
        elif f["status"] == "red" and f["value"] in (None, "", []):
            if f["type"] == "checkbox":
                edits.append({"field_id": f["field_id"], "value": True})
            elif f["type"] == "signature" or not f["required"]:
                edits.append({"field_id": f["field_id"], "acknowledged_blank": True})
            else:
                edits.append({"field_id": f["field_id"], "value": "Provided by researcher"})
    r = client.patch(f"/applications/{app_id}", json={"edits": edits})
    assert r.status_code == 200, r.text
    return r.json()


def test_thai_pdf_prefill_review_export(client, persona, grants_by_ref):
    grant = grants_by_ref["tedfund-young-startup"]
    tmpl = client.get(f"/grants/{grant.id}/form").json()
    assert tmpl["kind"] == "pdf_acroform" and len(tmpl["fields"]) == 33

    out = _create(client, persona, grant.id)
    app, summary = out["application"], out["summary"]
    by_id = {f["field_id"]: f for f in app["fields"]}

    # Pre-filled from the profile (green)
    assert by_id["pi_name"]["value"] == persona["applicant"]["name"] and by_id["pi_name"]["status"] == "green"
    assert by_id["trl"]["value"] == "4" and by_id["trl"]["status"] == "green"
    assert by_id["country"]["value"] == "Thailand"
    assert by_id["applicant_type"]["value"] == "University or research institute"
    assert by_id["use_pilot"]["value"] is True and by_id["use_personnel"]["value"] is False
    # Converted currency is amber, never silently green
    assert by_id["amount_thb"]["status"] == "amber"
    # Narrative fields are drafts with explicit placeholders - never fabricated specifics
    for fid in ("summary", "innovation", "objectives", "workplan", "impact", "budget_justification", "team"):
        assert by_id[fid]["narrative"] and PLACEHOLDER in by_id[fid]["value"] and by_id[fid]["status"] == "amber"
    # Declarations, signatures, dates, private keywords: always the researcher's job
    for fid in ("declaration", "applicant_signature", "date_signed", "keywords"):
        assert by_id[fid]["status"] == "red" and by_id[fid]["value"] in (None, "")
    filled = summary["green"] + summary["amber"]
    assert filled / summary["total"] >= 0.7, summary

    # Export is locked until review is complete - no skipping
    assert client.get(f"/applications/{app['id']}/export?format=pdf").status_code == 403
    r = client.post(f"/applications/{app['id']}/review", json={"attest": True})
    assert r.status_code == 422 and len(r.json()["problems"]) >= 10

    after = _resolve_everything(client, app["id"], app["fields"])
    assert after["problems"] == []
    assert client.post(f"/applications/{app['id']}/review", json={"attest": False}).status_code == 422
    assert client.post(f"/applications/{app['id']}/review", json={"attest": True}).status_code == 200

    pdf = client.get(f"/applications/{app['id']}/export?format=pdf")
    assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf"
    values = {k: v.get("/V") for k, v in PdfReader(io.BytesIO(pdf.content)).get_fields().items()}
    assert values["pi_name"] == persona["applicant"]["name"]
    assert values["use_pilot"] == "/Yes" and values["declaration"] == "/Yes"
    assert "Researcher-written text." in values["summary"]

    md = client.get(f"/applications/{app['id']}/export?format=md")
    assert md.status_code == 200 and "never submits" in md.text

    # Any edit re-opens review
    client.patch(f"/applications/{app['id']}", json={"edits": [{"field_id": "phone", "value": "+66 9"}]})
    assert client.get(f"/applications/{app['id']}/export?format=md").status_code == 403


def test_html_form_path(client, persona, grants_by_ref):
    grant = grants_by_ref["experiment-com"]
    out = _create(client, persona, grant.id)
    assert out["template"]["kind"] == "html_form"
    by_id = {f["field_id"]: f for f in out["application"]["fields"]}
    assert by_id["country"]["value"] == "Thailand"
    assert by_id["discipline"]["value"] == "Materials science"
    assert by_id["stage"]["value"] in ("Lab result", "Prototype") or by_id["stage"]["status"] == "red"
    assert set(by_id["budget_items"]["value"]) == {"Equipment"}
    assert by_id["accept_terms"]["status"] == "red"  # consent is never pre-ticked
    assert PLACEHOLDER in by_id["goals"]["value"]
    assert out["application"]["llm_used"] is False


def test_upload_flat_pdf_falls_back_to_layout(client, persona, grants_by_ref):
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = 780
    for line in ["Small Grant Application", "Name of applicant: ________", "Institution: ________",
                 "Email: ________", "Describe your project: ________", "☐ I agree to the terms"]:
        c.drawString(50, y, line)
        y -= 24
    c.save()
    r = client.post("/forms/upload", files={"file": ("flat.pdf", buf.getvalue(), "application/pdf")})
    assert r.status_code == 200, r.text
    t = r.json()
    assert t["kind"] == "pdf_flat" and t["extraction"] == "layout_heuristic"
    labels = [f["label"] for f in t["fields"]]
    assert "Name of applicant" in labels and "Institution" in labels

    out = _create(client, persona, grants_by_ref["nrct-research-grants"].id, template_id=t["id"])
    by_label = {f["label"]: f for f in out["application"]["fields"]}
    assert by_label["Institution"]["value"] == persona["applicant"]["organization"]
    # Flat PDFs can't be written into; export is the copy-paste pack
    app_id = out["application"]["id"]
    _resolve_everything(client, app_id, out["application"]["fields"])
    client.post(f"/applications/{app_id}/review", json={"attest": True})
    assert client.get(f"/applications/{app_id}/export?format=pdf").status_code == 400
    assert client.get(f"/applications/{app_id}/export?format=md").status_code == 200


def test_upload_rejects_non_forms(client):
    r = client.post("/forms/upload", files={"file": ("x.bin", b"\x00\x01binary", "application/octet-stream")})
    assert r.status_code == 400


def test_templates_are_cached_per_form(client, grants_by_ref):
    gid = grants_by_ref["nrct-research-grants"].id
    a = client.get(f"/grants/{gid}/form").json()
    b = client.get(f"/grants/{grants_by_ref['tedfund-young-startup'].id}/form").json()
    assert a["id"] == b["id"]  # same form_url -> one template, extracted once


DOCX_FORM = "thai-research-proposal-demo.docx"


def _docx_values(data: bytes) -> dict:
    from app.forms.docx_form import parse_docx

    return {s.id: s for s in parse_docx(data)[0]}


def test_thai_docx_form_prefill_review_export(client, persona, grants_by_ref):
    """Word forms: tables, content controls, legacy fields, dotted blanks and ☐ glyphs, all in Thai."""
    import docx

    from app.forms.fetch import DEMO_DIR

    raw = (DEMO_DIR / DOCX_FORM).read_bytes()
    r = client.post("/forms/upload", files={"file": (DOCX_FORM, raw, "application/octet-stream")})
    assert r.status_code == 200, r.text
    t = r.json()
    assert t["kind"] == "docx_form" and t["extraction"] == "docx" and len(t["fields"]) == 30
    ids = [f["id"] for f in t["fields"]]
    for fid in ("tbl0:r1:c1", "sdt:country", "ff:duration", "ff:declaration", "p7:cb3", "tbl1:r1:c0", "tbl2:r2:c0"):
        assert fid in ids, fid

    out = _create(client, persona, grants_by_ref["nrct-research-grants"].id, template_id=t["id"])
    app = out["application"]
    by_id = {f["field_id"]: f for f in app["fields"]}
    assert by_id["tbl0:r1:c1"]["value"] == persona["applicant"]["name"] and by_id["tbl0:r1:c1"]["status"] == "green"
    assert by_id["sdt:country"]["value"] == "ประเทศไทย"  # answered in the form's language
    assert by_id["p5:b0"]["value"] == "4"
    assert by_id["p5:b1"]["status"] == "amber"  # converted amount is never silently green
    assert by_id["p7:cb3"]["value"] is True and by_id["p7:cb2"]["value"] is False
    # Thai narrative cells get Thai drafts with Thai placeholders
    for fid in ("tbl1:r1:c0", "tbl1:r3:c0", "tbl1:r5:c0", "tbl1:r7:c0"):
        assert by_id[fid]["narrative"] and PLACEHOLDER_TH in by_id[fid]["value"] and by_id[fid]["status"] == "amber"
    # Declaration, signature, date and extra team rows stay with the researcher
    for fid in ("ff:declaration", "p12:b0", "p13:b0", "tbl2:r2:c0"):
        assert by_id[fid]["status"] == "red" and by_id[fid]["value"] in (None, "")

    app_id = app["id"]
    assert client.get(f"/applications/{app_id}/export?format=docx").status_code == 403
    after = _resolve_everything(client, app_id, app["fields"])
    assert after["problems"] == []
    assert client.post(f"/applications/{app_id}/review", json={"attest": True}).status_code == 200
    assert client.get(f"/applications/{app_id}/export?format=pdf").status_code == 400

    r = client.get(f"/applications/{app_id}/export?format=docx")
    assert r.status_code == 200 and r.headers["content-type"].endswith("wordprocessingml.document")
    filled = docx.Document(io.BytesIO(r.content))
    text = "\n".join(p.text for p in filled.paragraphs)
    cells = [c.text for tb in filled.tables for row in tb.rows for c in row.cells]
    assert persona["applicant"]["name"] in cells
    assert "☒ การผลิตทดลอง" in text and "☐ ค่าตอบแทนบุคลากร" in text
    assert "Researcher-written text." in "\n".join(cells)
    assert not any(has_placeholder(c) for c in cells)
    # Portal helper payload (tools/portal-fill): labelled answers, only after review
    portal = client.get(f"/applications/{app_id}/export?format=portal").json()
    assert portal["reviewed"] is True and portal["grant"]
    labels = {f["label"]: f for f in portal["fields"]}
    assert labels["อีเมล"]["value"] == persona["applicant"]["email"] and labels["อีเมล"]["type"] == "email"

    # Structural fields (content controls, legacy fields) survive filling and keep their values;
    # filled blanks and cells are no longer blanks, so they are not re-detected as empty fields.
    after = _docx_values(r.content)
    assert {"sdt:country", "ff:duration", "ff:declaration"} <= set(after)
    assert after["sdt:country"].label == "ประเทศ"


def test_docx_fill_unit():
    """Each field source writes where the reader expects, and nothing else changes."""
    import zipfile

    from app.forms.docx_form import fill_docx
    from app.forms.fetch import DEMO_DIR

    raw = (DEMO_DIR / DOCX_FORM).read_bytes()
    out = fill_docx(raw, {"sdt:country": "ประเทศไทย", "ff:duration": "9", "ff:declaration": True,
                          "p4:b0": "วัสดุศาสตร์", "p7:cb0": True, "tbl0:r0:c1": "Test project"})
    xml = zipfile.ZipFile(io.BytesIO(out)).read("word/document.xml").decode()
    assert "ประเทศไทย" in xml and ">9<" in xml and "Test project" in xml
    assert "☒ ครุภัณฑ์" in xml or "☒</w:t>" in xml
    assert 'w:default w:val="1"' in xml or "w:checked" in xml  # legacy checkbox ticked
    assert "สาขาการวิจัย วัสดุศาสตร์" in xml or "วัสดุศาสตร์" in xml
    # Untouched blanks stay blank
    assert "งบประมาณที่ขอรับการสนับสนุน (บาท) ...." in xml
    # Values the caller did not send are not invented
    unchanged = fill_docx(raw, {})
    assert _docx_values(unchanged).keys() == _docx_values(raw).keys()
