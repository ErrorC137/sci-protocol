"""Word-form parser/filler safety: a wrong label is worse than no label, and bad files are 400s, not 500s.

Each test builds a small .docx in memory from raw OOXML so the exact structure under test is explicit."""

import io
import zipfile

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from app.forms.docx_form import fill_docx, parse_docx

NS = nsdecls("w") + ' xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"'


# -- tiny OOXML builders ---------------------------------------------------------------------------------------
def r(text: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def p(*parts: str) -> str:
    return f"<w:p>{''.join(parts)}</w:p>"


def sdt(content: str, pr: str = "") -> str:
    return f"<w:sdt><w:sdtPr>{pr}</w:sdtPr><w:sdtContent>{content}</w:sdtContent></w:sdt>"


def tag(v: str) -> str:
    return f'<w:tag w:val="{v}"/>'


CHECKBOX_PR = ('<w14:checkbox><w14:checkedState w14:val="2612" w14:font="MS Gothic"/>'
               '<w14:uncheckedState w14:val="2610" w14:font="MS Gothic"/></w14:checkbox>')


def tc(content: str = "", span: int = 1, vmerge: str | None = None) -> str:
    pr = (f'<w:gridSpan w:val="{span}"/>' if span > 1 else "") + (
        '<w:vMerge w:val="restart"/>' if vmerge == "restart" else "<w:vMerge/>" if vmerge else "")
    return f"<w:tc><w:tcPr>{pr}</w:tcPr>{content or '<w:p/>'}</w:tc>"


def tr(*cells: str) -> str:
    return f"<w:tr>{''.join(cells)}</w:tr>"


def tbl(*rows: str) -> str:
    return f"<w:tbl><w:tblPr/><w:tblGrid/>{''.join(rows)}</w:tbl>"


def _fld(kind: str) -> str:
    return f'<w:r><w:fldChar w:fldCharType="{kind}"/></w:r>'


def ff_text(name: str | None, result: str | None = None, separate: bool = True) -> str:
    nm = f'<w:name w:val="{name}"/>' if name else ""
    return (f'<w:r><w:rPr><w:b/></w:rPr><w:fldChar w:fldCharType="begin"><w:ffData>{nm}<w:enabled/><w:textInput/>'
            '</w:ffData></w:fldChar></w:r><w:r><w:instrText xml:space="preserve"> FORMTEXT </w:instrText></w:r>'
            + (_fld("separate") if separate else "") + (r(result) if result is not None else "") + _fld("end"))


def ff_box(name: str) -> str:
    return (f'<w:r><w:fldChar w:fldCharType="begin"><w:ffData><w:name w:val="{name}"/><w:enabled/><w:checkBox>'
            '<w:sizeAuto/><w:default w:val="0"/></w:checkBox></w:ffData></w:fldChar></w:r>'
            '<w:r><w:instrText xml:space="preserve"> FORMCHECKBOX </w:instrText></w:r>' + _fld("end"))


def ff_dropdown(name: str, entries: list[str]) -> str:
    items = "".join(f'<w:listEntry w:val="{e}"/>' for e in entries)
    return (f'<w:r><w:fldChar w:fldCharType="begin"><w:ffData><w:name w:val="{name}"/><w:enabled/><w:ddList>{items}'
            '</w:ddList></w:ffData></w:fldChar></w:r>'
            '<w:r><w:instrText xml:space="preserve"> FORMDROPDOWN </w:instrText></w:r>' + _fld("separate") + _fld("end"))


def docx(*blocks: str) -> bytes:
    d = Document()
    body = d.element.body
    sect = body.find(qn("w:sectPr"))
    for el in list(parse_xml(f"<w:body {NS}>{''.join(blocks)}</w:body>")):
        sect.addprevious(el)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def specs(data: bytes) -> dict:
    return {s.id: s for s in parse_docx(data)[0]}


def doc_xml(data: bytes) -> str:
    return zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode()


def no_invented_labels(found: dict) -> None:
    assert not any(s.label.startswith("Field ") for s in found.values())


# -- 1. content controls without alias/tag ---------------------------------------------------------------------
def test_sdt_caption_from_same_paragraph_before_the_control():
    found = specs(docx(p(r("Name: "), sdt(r("Click here")), r("  Position: "), sdt(r("Click here"))),
                       p(r("งบประมาณที่ขอ "), sdt(r("0")), r(" บาท"))))
    assert [s.label for s in found.values()] == ["Name", "Position", "งบประมาณที่ขอ (บาท)"]
    no_invented_labels(found)


def test_sdt_caption_from_left_cell_then_cell_above():
    found = specs(docx(tbl(tr(tc(p(r("Applicant name"))), tc(sdt(p(r("Click here")))))),
                       tbl(tr(tc(), tc(p(r("Project summary")))), tr(tc(), tc(sdt(p(r("Click here"))))))))
    assert [s.label for s in found.values()] == ["Applicant name", "Project summary"]


def test_sdt_caption_from_previous_paragraph_or_empty():
    found = specs(docx(sdt(p(r("Click here"))), p(r("Project summary")), p(), sdt(p(r("Click here")))))
    assert found["sdt:0"].label == ""  # nothing to go on: empty, never "Field 1"
    assert found["sdt:1"].label == "Project summary"
    no_invented_labels(found)


def test_checkbox_sdt_caption_after_box_and_fill_creates_checked_element():
    data = docx(p(sdt(r("☐"), CHECKBOX_PR), r(" I agree to share my data")))
    found = specs(data)
    assert found["sdt:0"].label == "I agree to share my data" and found["sdt:0"].type == "checkbox"
    xml = doc_xml(fill_docx(data, {"sdt:0": True}))
    assert '<w14:checked w14:val="1"/>' in xml and "☒" in xml and "☐" not in xml


# -- 2. duplicate tags ------------------------------------------------------------------------------------------
def test_duplicate_sdt_tags_get_unique_ids_that_fill_the_right_control():
    data = docx(*(p(r(f"{k}: "), sdt(r(f"old{k}"), tag("person"))) for k in "ABC"))
    assert list(specs(data)) == ["sdt:person", "sdt:person:1", "sdt:person:2"]
    out = fill_docx(data, {"sdt:person": "Alice", "sdt:person:1": "Bob", "sdt:person:2": "Carol"})
    xml = doc_xml(out)
    assert xml.index("Alice") < xml.index("Bob") < xml.index("Carol") and "old" not in xml
    assert list(specs(out)) == ["sdt:person", "sdt:person:1", "sdt:person:2"]


# -- 3. row / table level content controls ----------------------------------------------------------------------
def test_table_and_row_level_sdts_are_containers_not_fields():
    found = specs(docx(
        sdt(tbl(tr(tc(p(r("Institution"))), tc()))),                                   # whole table in a block sdt
        tbl(tr(tc(p(r("Name"))), tc(p(r("Position")))), sdt(tr(tc(), tc()), tag("rows"))),  # repeating-row sdt
        tbl(tr(tc(p(r("Email"))), sdt(tc(), tag("cell")))),                             # cell-level sdt
    ))
    assert not any(k.startswith("sdt:") for k in found)
    assert found["tbl0:r0:c1"].label == "Institution"
    assert found["tbl1:r1:c0"].label == "Name" and found["tbl1:r1:c1"].label == "Position"
    assert found["tbl2:r0:c1"].label == "Email"


# -- 4. legacy form fields --------------------------------------------------------------------------------------
def test_formtext_caption_excludes_other_fields_values():
    found = specs(docx(p(r("Name: "), ff_text("pi_name", "Somchai Jaidee"), r("  Position: "), ff_text("pi_pos", ""))))
    assert found["ff:pi_name"].label == "Name" and found["ff:pi_pos"].label == "Position"


def test_formcheckbox_caption_is_text_after_the_box():
    found = specs(docx(p(ff_box("c1"), r(" Yes  "), ff_box("c2"), r(" No")),
                       p(r("Yes "), ff_box("c3"), r("  No "), ff_box("c4"))))  # captions-first layout
    assert [found[f"ff:c{k}"].label for k in range(1, 5)] == ["Yes", "No", "Yes", "No"]


def test_form_field_without_name_or_caption_has_empty_label():
    found = specs(docx(p(ff_text(None)), p(ff_text("Text1"))))
    assert found["ff:ff0"].label == ""
    assert found["ff:Text1"].label == "Text1"  # named: falls back to the humanized name


def _field_runs(data: bytes, para_idx: int) -> list[str]:
    para = Document(io.BytesIO(data)).element.body.findall(qn("w:p"))[para_idx]
    out = []
    for x in para.iter(qn("w:fldChar"), qn("w:t")):
        out.append(x.get(qn("w:fldCharType")) if x.tag == qn("w:fldChar") else x.text)
    return out


def test_fill_formtext_with_empty_result_inserts_the_value():
    data = docx(p(r("Name: "), ff_text("a")), p(r("Org: "), ff_text("b", separate=False)))
    out = fill_docx(data, {"ff:a": "Alice", "ff:b": "Acme"})
    assert _field_runs(out, 0) == ["Name: ", "begin", "separate", "Alice", "end"]
    assert _field_runs(out, 1) == ["Org: ", "begin", "separate", "Acme", "end"]
    para = Document(io.BytesIO(out)).element.body.findall(qn("w:p"))[0]
    value_run = next(t for t in para.iter(qn("w:t")) if t.text == "Alice").getparent()
    assert value_run.find(qn("w:rPr")).find(qn("w:b")) is not None  # formatting copied from the begin run
    found = specs(out)
    assert found["ff:a"].label == "Name" and found["ff:b"].label == "Org"


def test_fill_formdropdown_sets_selected_index_only_for_real_options():
    data = docx(p(r("Country: "), ff_dropdown("country", ["Laos", "Thailand", "Vietnam"])))
    assert specs(data)["ff:country"].options == ["Laos", "Thailand", "Vietnam"]
    dd = Document(io.BytesIO(fill_docx(data, {"ff:country": "Thailand"}))).element.body.find(".//" + qn("w:ddList"))
    assert dd[0].tag == qn("w:result") and dd[0].get(qn("w:val")) == "1"
    dd = Document(io.BytesIO(fill_docx(data, {"ff:country": "Japan"}))).element.body.find(".//" + qn("w:ddList"))
    assert dd.find(qn("w:result")) is None


# -- 5. merged table cells --------------------------------------------------------------------------------------
def test_gridspan_header_lines_up_by_logical_column():
    found = specs(docx(tbl(tr(tc(p(r("Item")), span=2), tc(p(r("Amount (บาท)")))), tr(tc(), tc(), tc()))))
    assert found["tbl0:r1:c2"].label == "Amount (บาท)"
    assert found["tbl0:r1:c1"].label == "Item"  # under the spanned header, not the next physical cell


def test_vmerge_continuations_are_never_fields_and_resolve_for_labels_above():
    found = specs(docx(
        tbl(tr(tc(p(r("Address")), vmerge="restart"), tc()), tr(tc(vmerge="continue"), tc())),
        tbl(tr(tc(p(r("Name")), vmerge="restart")), tr(tc(vmerge="continue")), tr(tc())),
    ))
    assert found["tbl0:r0:c1"].label == "Address"
    assert "tbl0:r1:c0" not in found and "tbl1:r1:c0" not in found
    assert found["tbl1:r2:c0"].label == "Name"


# -- 6. content-control text is opaque --------------------------------------------------------------------------
def test_control_text_is_never_a_blank_or_a_caption():
    found = specs(docx(
        p(r("Name: "), sdt(r(".........."))),                          # dots inside a control are not a blank
        tbl(tr(tc(sdt(p(r("Applicant")))), tc())),                     # a control's text does not label the next cell
        p(sdt(r("Applicant name"))), sdt(p(r("Click here"))),          # ... nor the control below it
    ))
    assert "p0:b0" not in found and found["sdt:0"].label == "Name"
    assert "tbl0:r0:c1" not in found
    assert found["sdt:3"].label == ""


# -- 7. paragraph blanks ----------------------------------------------------------------------------------------
def test_blank_caption_stops_at_checkbox_and_takes_unit():
    found = specs(docx(
        p(r("ข้อมูลผู้สมัคร")),
        p(r("☐ ครุภัณฑ์ ☐ อื่นๆ ระบุ ..........")),
        p(r("งบประมาณที่ขอ .......... บาท ระยะเวลา .......... เดือน")),
        p(r("เริ่ม .......... วันที่ ..........")),
    ))
    assert found["p1:b0"].label == "อื่นๆ ระบุ"
    assert found["p1:cb0"].label == "ครุภัณฑ์" and found["p1:cb1"].label == "อื่นๆ ระบุ"
    assert found["p2:b0"].label == "งบประมาณที่ขอ (บาท)" and found["p2:b1"].label == "ระยะเวลา (เดือน)"
    assert found["p3:b0"].label == "เริ่ม" and found["p3:b1"].label == "วันที่"  # "วันที่" is not the unit "วัน"
    assert found["p2:b0"].hint == "งบประมาณที่ขอ … บาท ระยะเวลา … เดือน | prev: ☐ ครุภัณฑ์ ☐ อื่นๆ ระบุ …"


# -- 8. label-less dotted lines under a heading -----------------------------------------------------------------
def test_dotted_lines_under_heading_merge_into_one_field():
    line = p(r("." * 60))
    data = docx(p(r("บทคัดย่อ")), line, line, p(), line, p(r("ที่อยู่ ..........")), line)
    found = specs(data)
    assert found["p1:b0"].label == "บทคัดย่อ" and found["p1:b0"].narrative and found["p1:b0"].type == "textarea"
    assert not {"p2:b0", "p4:b0", "p6:b0"} & set(found)  # merged; an address continuation line is not a new field
    assert found["p5:b0"].label == "ที่อยู่"
    paras = Document(io.BytesIO(fill_docx(data, {"p1:b0": "บรรทัดแรก\nบรรทัดสอง"}))).paragraphs
    assert paras[1].text == "บรรทัดแรก\nบรรทัดสอง" and paras[2].text == "" and paras[4].text == ""
    assert paras[6].text == "." * 60


# -- 9. data rows of repeating tables ---------------------------------------------------------------------------
def test_table_data_rows_carry_table_row_hint():
    found = specs(docx(
        p(r("ผลงานตีพิมพ์")),
        tbl(tr(tc(p(r("ชื่อเรื่อง"))), tc(p(r("ปีที่พิมพ์")))), tr(tc(), tc()), tr(tc(), tc())),
        p(r("ข้อมูลติดต่อ")),
        tbl(tr(tc(p(r("อีเมล"))), tc())),
        p(r("คณะผู้วิจัย")),
        tbl(tr(tc(p(r("ลำดับ"))), tc(p(r("ชื่อ-สกุล")))), tr(tc(p(r("1"))), tc()), tr(tc(p(r("2"))), tc())),
    ))
    hint = "[table row] ผลงานตีพิมพ์ | ชื่อเรื่อง / ปีที่พิมพ์"
    assert found["tbl0:r1:c0"].label == "ชื่อเรื่อง" and found["tbl0:r1:c0"].hint == hint
    assert found["tbl0:r2:c1"].label == "ปีที่พิมพ์ (2)" and found["tbl0:r2:c1"].hint == hint
    assert found["tbl1:r0:c1"].label == "อีเมล" and found["tbl1:r0:c1"].hint is None
    assert found["tbl2:r1:c1"].label == "ชื่อ-สกุล"  # the row number "1" is not a caption
    assert found["tbl2:r2:c1"].label == "ชื่อ-สกุล (2)"
    assert found["tbl2:r1:c1"].hint == "[table row] คณะผู้วิจัย | ลำดับ / ชื่อ-สกุล"


def test_label_above_narrative_cell_is_not_a_table_row():
    found = specs(docx(tbl(tr(tc(p(r("วัตถุประสงค์")))), tr(tc()), tr(tc(p(r("แผนการดำเนินงาน")))), tr(tc()))))
    assert found["tbl0:r1:c0"].hint is None and found["tbl0:r3:c0"].hint is None


# -- 10. zip bombs ----------------------------------------------------------------------------------------------
def _zip(entries: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED, compresslevel=1) as z:
        for name, blob in entries.items():
            z.writestr(name, blob)
    return buf.getvalue()


@pytest.mark.parametrize("entries, message", [
    ({"word/document.xml": b" " * (31 << 20)}, "too large"),
    ({"word/document.xml": b"<x/>", "word/media/big.bin": b"\0" * (101 << 20)}, "too large when unpacked"),
    ({**{f"word/media/{k}.bin": b"" for k in range(2001)}, "word/document.xml": b"<x/>"}, "too many parts"),
], ids=["document.xml>30MB", "total>100MB", "entries>2000"])
def test_zip_bomb_guard(entries, message):
    data = _zip(entries)
    assert len(data) < 1 << 20  # small on the wire, huge when inflated
    for fn in (parse_docx, lambda d: fill_docx(d, {})):
        with pytest.raises(ValueError, match=message):
            fn(data)


# -- 11. malformed files are ValueErrors (HTTP 400) -------------------------------------------------------------
def _broken_docx(document_xml: bytes) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(docx(p(r("x")))))
    return _zip({n: (document_xml if n == "word/document.xml" else src.read(n)) for n in src.namelist()})


@pytest.mark.parametrize("data", [
    b"not a zip at all",
    b"PK\x03\x04truncated",
    _zip({"hello.txt": b"hi"}),
    _broken_docx(b"<w:document"),
    _broken_docx(f"<w:document {nsdecls('w')}/>".encode()),
], ids=["not-zip", "truncated-zip", "zip-without-docx-parts", "malformed-xml", "no-body"])
def test_unreadable_files_raise_value_error(data):
    with pytest.raises(ValueError):
        parse_docx(data)
    with pytest.raises(ValueError):
        fill_docx(data, {"p0:b0": "x"})


def test_upload_of_corrupt_docx_is_a_400(client):
    r_ = client.post("/forms/upload", files={"file": ("bad.docx", _broken_docx(b"<w:document"), "application/octet-stream")})
    assert r_.status_code == 400 and "Word" in r_.json()["detail"]


# -- 12. values are sanitised before they reach lxml ------------------------------------------------------------
def test_fill_strips_xml_illegal_characters():
    data = docx(p(r("ชื่อ ..........")), tbl(tr(tc(p(r("อีเมล"))), tc())), p(r("ตำแหน่ง "), ff_text("pos", "")),
                p(r("หน่วยงาน: "), sdt(r("x"))))
    dirty = "A\x00B\x0bC\x1fD\ud800E\ufffeF"
    out = fill_docx(data, {"p0:b0": dirty, "tbl0:r0:c1": dirty, "ff:pos": dirty, "sdt:0": [dirty, "G\x08"]})
    xml = doc_xml(out)
    assert xml.count("ABCDEF") == 4 and "ABCDEF, G" in xml
