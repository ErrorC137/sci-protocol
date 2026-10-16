"""Word (.docx) application forms - the most common format for Thai agency forms.

Field sources, in document order:
  sdt   content controls (w:sdt): text, checkbox, dropdown, date
  ff    legacy form fields (FORMTEXT / FORMCHECKBOX / FORMDROPDOWN)
  cell  a table cell that is blank (or only dots/underscores) next to, or below, a label cell
  blank a run of dots/underscores in a paragraph ("ชื่อ.................. นามสกุล..................")
  cb    a ☐/□ glyph followed by its label text

parse() and fill() walk the document identically, so field ids are stable between the two.
"""

from __future__ import annotations

import copy
import io
import re
from dataclasses import dataclass
from typing import Callable

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from ..models.schemas import FieldSpec
from .fieldkind import guess_type, humanize, is_narrative

W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
BLANK_RE = re.compile(r"[\.…_]{4,}|(?:\. ){4,}")
CHECK_EMPTY, CHECK_FULL = "☐□", "☒"
ONLY_BLANK = re.compile(r"^[\s\.…_:：]*$")
MAX_LABEL = 160


def _text(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _clean_label(s: str) -> str:
    s = BLANK_RE.sub(" ", s)
    s = re.sub(r"\s+", " ", s).strip(" :：-–\t")
    return s[:MAX_LABEL]


@dataclass
class _Site:
    """Where a field lives in the XML, so fill() can write to it."""

    kind: str
    el: object
    extra: dict


def _walk(doc) -> list[tuple[FieldSpec, _Site]]:
    body = doc.element.body
    found: list[tuple[FieldSpec, _Site]] = []
    seen_sdt = set()

    # 1. content controls -------------------------------------------------------
    for i, sdt in enumerate(body.iter(qn("w:sdt"))):
        pr = sdt.find(qn("w:sdtPr"))
        if pr is None:
            continue
        alias = pr.find(qn("w:alias"))
        tag = pr.find(qn("w:tag"))
        label = (alias.get(qn("w:val")) if alias is not None else None) or (tag.get(qn("w:val")) if tag is not None else None)
        content = sdt.find(qn("w:sdtContent"))
        if not label:
            prev = sdt.getprevious()
            label = _clean_label(_text(prev)) if prev is not None else f"Field {i + 1}"
        fid = f"sdt:{(tag.get(qn('w:val')) if tag is not None else None) or i}"
        ftype, options = "text", []
        if pr.find(f"{{{W14}}}checkbox") is not None:
            ftype = "checkbox"
        elif pr.find(qn("w:dropDownList")) is not None or pr.find(qn("w:comboBox")) is not None:
            lst = pr.find(qn("w:dropDownList")) if pr.find(qn("w:dropDownList")) is not None else pr.find(qn("w:comboBox"))
            ftype = "select"
            options = [li.get(qn("w:displayText")) or li.get(qn("w:value")) for li in lst.findall(qn("w:listItem"))]
        elif pr.find(qn("w:date")) is not None:
            ftype = "date"
        ftype = guess_type(label, ftype)
        seen_sdt.add(id(sdt))
        found.append((FieldSpec(id=fid, label=label, type=ftype, options=options, narrative=is_narrative(label, ftype, None)),
                      _Site("sdt", content, {"sdt": sdt})))

    # 2. legacy form fields -------------------------------------------------------
    for i, fc in enumerate(body.iter(qn("w:fldChar"))):
        if fc.get(qn("w:fldCharType")) != "begin":
            continue
        ff = fc.find(qn("w:ffData"))
        if ff is None:
            continue
        name_el = ff.find(qn("w:name"))
        name = name_el.get(qn("w:val")) if name_el is not None else f"ff{i}"
        run = fc.getparent()
        para = run.getparent()
        label = _clean_label(_text(para).split("FORMTEXT")[0]) or humanize(name)
        if ff.find(qn("w:checkBox")) is not None:
            ftype, maxlen, options = "checkbox", None, []
        elif ff.find(qn("w:ddList")) is not None:
            ftype, maxlen = "select", None
            options = [e.get(qn("w:val")) for e in ff.find(qn("w:ddList")).findall(qn("w:listEntry"))]
        else:
            ti = ff.find(qn("w:textInput"))
            ml = ti.find(qn("w:maxLength")) if ti is not None else None
            ftype, maxlen, options = "text", (int(ml.get(qn("w:val"))) if ml is not None else None), []
        ftype = guess_type(label, ftype)
        found.append((FieldSpec(id=f"ff:{name}", label=label, type=ftype, options=options, maxlen=maxlen,
                                narrative=is_narrative(label, ftype, maxlen)), _Site("ff", run, {"ff": ff})))

    # 3/4/5. tables, paragraph blanks, glyph checkboxes - in body order ----------------
    def in_sdt(el) -> bool:
        p = el.getparent()
        while p is not None:
            if p.tag == qn("w:sdt"):
                return True
            p = p.getparent()
        return False

    def para_fields(p, prefix: str) -> list[tuple[FieldSpec, _Site]]:
        if in_sdt(p) or any(True for _ in p.iter(qn("w:fldChar"))):
            return []
        text = _text(p)
        out = []
        for k, m in enumerate(BLANK_RE.finditer(text)):
            before = text[: m.start()]
            prev_blank = list(BLANK_RE.finditer(before))
            label = _clean_label(before[prev_blank[-1].end():] if prev_blank else before)
            if not label:
                continue
            ftype = guess_type(label)
            out.append((FieldSpec(id=f"{prefix}:b{k}", label=label, type=ftype, narrative=is_narrative(label, ftype, None)),
                        _Site("blank", p, {"start": m.start(), "end": m.end()})))
        for k, m in enumerate(re.finditer(f"[{CHECK_EMPTY}]", text)):
            rest = text[m.end():]
            nxt = re.search(f"[{CHECK_EMPTY}{CHECK_FULL}]", rest)
            label = _clean_label(rest[: nxt.start()] if nxt else rest)
            if label:
                out.append((FieldSpec(id=f"{prefix}:cb{k}", label=label, type="checkbox"), _Site("glyph", p, {"pos": m.start()})))
        return out

    for t_idx, tbl in enumerate(body.iter(qn("w:tbl"))):
        rows = tbl.findall(qn("w:tr"))
        grid = [[c for c in r.findall(qn("w:tc"))] for r in rows]
        for r, cells in enumerate(grid):
            for c, cell in enumerate(cells):
                ctext = _text(cell)
                if in_sdt(cell) or any(True for _ in cell.iter(qn("w:fldChar"))):
                    continue
                if not ONLY_BLANK.match(ctext):
                    # label cells may still hold inline blanks/checkboxes
                    for i, p in enumerate(cell.findall(qn("w:p"))):
                        found.extend(para_fields(p, f"tbl{t_idx}:r{r}:c{c}:p{i}"))
                    continue
                label = ""
                if c > 0 and not ONLY_BLANK.match(_text(cells[c - 1])):
                    label = _clean_label(_text(cells[c - 1]))
                elif r > 0 and c < len(grid[r - 1]) and not ONLY_BLANK.match(_text(grid[r - 1][c])):
                    label = _clean_label(_text(grid[r - 1][c]))
                else:  # data rows of a multi-row table: use the column header, numbered by row
                    for up in range(r - 1, -1, -1):
                        if c < len(grid[up]) and not ONLY_BLANK.match(_text(grid[up][c])):
                            label = f"{_clean_label(_text(grid[up][c]))} ({r - up})"
                            break
                if not label:
                    continue
                big = len(cell.findall(qn("w:p"))) >= 3 or (r > 0 and label == _clean_label(_text(grid[r - 1][c])))
                ftype = "textarea" if big and is_narrative(label, "text", None) else guess_type(label)
                found.append((FieldSpec(id=f"tbl{t_idx}:r{r}:c{c}", label=label, type=ftype,
                                        narrative=is_narrative(label, ftype, None)), _Site("cell", cell, {})))

    for p_idx, p in enumerate(body.findall(qn("w:p"))):
        found.extend(para_fields(p, f"p{p_idx}"))

    # Present fields in reading order (document position, then offset within a paragraph).
    order = {id(el): i for i, el in enumerate(body.iter())}
    found.sort(key=lambda fs: (order.get(id(fs[1].el), 0), fs[1].extra.get("start", fs[1].extra.get("pos", 0))))
    return found


def parse_docx(data: bytes) -> tuple[list[FieldSpec], str | None]:
    doc = Document(io.BytesIO(data))
    title = doc.core_properties.title or None
    if not title:
        first = next((p.text.strip() for p in doc.paragraphs if p.text.strip()), None)
        title = first[:120] if first else None
    specs, seen = [], set()
    for spec, _ in _walk(doc):
        if spec.id not in seen:
            seen.add(spec.id)
            specs.append(spec)
    return specs, title


# ---------------------------------------------------------------------------
# Filling
# ---------------------------------------------------------------------------

def _as_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, list):
        return ", ".join(v)
    if isinstance(v, bool):
        return "Yes" if v else "No"
    return str(v)


def _replace_span(p, start: int, end: int, new: str) -> None:
    """Replace characters [start, end) of a paragraph's text across runs, keeping the first run's formatting."""
    pos = 0
    first_t = None
    for t in p.iter(qn("w:t")):
        s = t.text or ""
        a, b = pos, pos + len(s)
        pos = b
        if b <= start or a >= end:
            continue
        lo, hi = max(start, a) - a, min(end, b) - a
        if first_t is None:
            first_t = t
            t.text = s[:lo] + new + s[hi:]
        else:
            t.text = s[:lo] + s[hi:]
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def _set_runs_text(container, value: str, template_run=None) -> None:
    """Make the container's (cell / sdtContent) first paragraph contain exactly `value`."""
    paras = container.findall(qn("w:p")) or [container]
    p = paras[0]
    runs = p.findall(qn("w:r"))
    base = template_run if template_run is not None else (runs[0] if runs else None)
    for r in runs:
        p.remove(r)
    new_r = copy.deepcopy(base) if base is not None else etree.SubElement(p, qn("w:r"))
    for child in list(new_r):
        if child.tag != qn("w:rPr"):
            new_r.remove(child)
    lines = value.split("\n")
    for i, line in enumerate(lines):
        t = etree.SubElement(new_r, qn("w:t"))
        t.text = line
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        if i < len(lines) - 1:
            etree.SubElement(new_r, qn("w:br"))
    if new_r.getparent() is None:
        p.append(new_r)


def fill_docx(data: bytes, values: dict[str, object]) -> bytes:
    doc = Document(io.BytesIO(data))
    sites = _walk(doc)
    # Paragraph blanks shift character offsets, so fill them right-to-left within each paragraph.
    sites.sort(key=lambda fs: -(fs[1].extra.get("start", fs[1].extra.get("pos", 0))))

    def do_sdt(spec, site, v):
        if spec.type == "checkbox":
            box = site.extra["sdt"].find(qn("w:sdtPr")).find(f"{{{W14}}}checkbox")
            checked = box.find(f"{{{W14}}}checked")
            if checked is not None:
                checked.set(f"{{{W14}}}val", "1" if v is True else "0")
            for t in site.el.iter(qn("w:t")):
                t.text = CHECK_FULL if v is True else "☐"
            return
        _set_runs_text(site.el, _as_text(v))

    def do_ff(spec, site, v):
        ff = site.extra["ff"]
        if spec.type == "checkbox":
            cb = ff.find(qn("w:checkBox"))
            checked = cb.find(qn("w:checked"))
            if checked is None:
                checked = etree.SubElement(cb, qn("w:checked"))
            checked.set(qn("w:val"), "1" if v is True else "0")
            return
        # value goes in the first text run between the 'separate' and 'end' field characters
        run = site.el
        sib = run.getnext()
        in_result = False
        while sib is not None:
            fc = sib.find(qn("w:fldChar"))
            if fc is not None and fc.get(qn("w:fldCharType")) == "separate":
                in_result = True
            elif fc is not None and fc.get(qn("w:fldCharType")) == "end":
                break
            elif in_result:
                ts = sib.findall(qn("w:t"))
                if ts:
                    ts[0].text = _as_text(v)
                    for extra in ts[1:]:
                        extra.text = ""
                    return
            sib = sib.getnext()

    def do_cell(spec, site, v):
        _set_runs_text(site.el, _as_text(v))

    def do_blank(spec, site, v):
        _replace_span(site.el, site.extra["start"], site.extra["end"], f" {_as_text(v)} ")

    def do_glyph(spec, site, v):
        if v is True:
            _replace_span(site.el, site.extra["pos"], site.extra["pos"] + 1, CHECK_FULL)

    handlers: dict[str, Callable] = {"sdt": do_sdt, "ff": do_ff, "cell": do_cell, "blank": do_blank, "glyph": do_glyph}
    done = set()
    for spec, site in sites:
        if spec.id in done or spec.id not in values:
            continue
        v = values[spec.id]
        if v in (None, "", []):
            continue
        handlers[site.kind](spec, site, v)
        done.add(spec.id)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
