"""Fetched bytes -> clean text + links. HTML via trafilatura (boilerplate removed, tables kept),
PDF via pdfplumber, Word via python-docx.

Thai government sites still serve TIS-620 / windows-874 pages, sometimes with a wrong or missing
charset header; bs4's UnicodeDammit sniffs the real encoding instead of producing mojibake.
"""

from __future__ import annotations

import hashlib
import io
import logging
import re
from dataclasses import dataclass, field
from typing import TYPE_CHECKING
from urllib.parse import urljoin

from bs4 import BeautifulSoup, UnicodeDammit

from .. import thai

if TYPE_CHECKING:
    from .crawl import Page

log = logging.getLogger("sci.curate")

MAX_PDF_PAGES = 40


@dataclass
class Doc:
    url: str
    kind: str  # html | pdf | docx | other
    title: str | None
    text: str
    links: list[tuple[str, str]] = field(default_factory=list)  # (absolute url, anchor text)

    @property
    def sha(self) -> str:
        """Hash of the extracted text, not the bytes: a rotated banner or CSRF token is not a change."""
        return hashlib.sha256(self.text.encode()).hexdigest()

    @property
    def lang(self) -> str:
        return "th" if thai.is_thai(self.text) else "en"


def _clean(text: str) -> str:
    text = thai.normalize(text)
    text = re.sub(r"[ \t]*\n[ \t]*", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


# Charsets servers announce by default whatever the page really is (Apache/IIS defaults). Treated as no hint.
_WEAK_CHARSETS = {"iso-8859-1", "latin-1", "latin1", "us-ascii", "ascii", "windows-1252", "cp1252"}
_MOJIBAKE = re.compile(r"[\u00c0-\u00ff]{3,}")


def decode_html(data: bytes, content_type: str = "") -> str:
    """Header charset is a hint, not the truth: Thai government servers often send ISO-8859-1 for
    windows-874 pages. A <meta> charset and byte sniffing win over a weak header, and Latin-1-looking
    mojibake is re-decoded as cp874 when that yields Thai."""
    m = re.search(r"charset=([\w-]+)", content_type or "", re.I)
    header = m.group(1).lower() if m else None
    hints = [header] if header and header not in _WEAK_CHARSETS else []
    dammit = UnicodeDammit(data, known_definite_encodings=hints, smart_quotes_to="html", is_html=True)
    text = dammit.unicode_markup or data.decode("utf-8", "replace")
    if len(_MOJIBAKE.findall(text)) >= 3:
        try:
            thai_text = data.decode("cp874")
            if thai.is_thai(thai_text):
                return thai_text
        except UnicodeDecodeError:
            pass
    return text


def html_doc(url: str, data: bytes, content_type: str = "") -> Doc:
    html = decode_html(data, content_type)
    soup = BeautifulSoup(html, "lxml")
    og = soup.find("meta", property="og:title")
    h1 = soup.find("h1")
    title = (og.get("content") if og else None) or (soup.title.get_text(" ", strip=True) if soup.title else None) \
        or (h1.get_text(" ", strip=True) if h1 else None)
    links = []
    for a in soup.find_all("a", href=True):
        href = a["href"].strip()
        if href.startswith(("mailto:", "tel:", "javascript:", "#")):
            continue
        try:
            links.append((urljoin(url, href), a.get_text(" ", strip=True)))
        except ValueError:  # 'http://[broken/x' and similar: drop the link, keep the page
            continue
    text = None
    try:
        import trafilatura

        text = trafilatura.extract(html, url=url, favor_recall=True, include_tables=True, include_comments=False)
    except Exception as ex:  # trafilatura is an optimisation; plain text is the fallback
        log.debug("trafilatura failed on %s: %s", url, ex)
    if not text or len(text) < 200:
        for tag in soup(["script", "style", "noscript", "nav", "footer"]):
            tag.decompose()
        text = soup.get_text("\n", strip=True)
    return Doc(url, "html", _clean(title) if title else None, _clean(text), links)


def pdf_doc(url: str, data: bytes) -> Doc:
    import pdfplumber

    parts = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages[:MAX_PDF_PAGES]:
            parts.append(page.extract_text() or "")
    text = _clean("\n".join(parts))
    first = next((ln for ln in text.splitlines() if len(ln.strip()) > 8), None)
    return Doc(url, "pdf", first, text, [])


def docx_doc(url: str, data: bytes) -> Doc:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    text = _clean("\n".join(parts))
    first = next((ln for ln in text.splitlines() if len(ln.strip()) > 8), None)
    return Doc(url, "docx", first, text, [])


def textify(page: Page) -> Doc:
    ct, data, url = page.content_type.lower(), page.data, page.final_url
    if "pdf" in ct or data[:5] == b"%PDF-":
        return pdf_doc(url, data)
    if "wordprocessingml" in ct or (data[:4] == b"PK\x03\x04" and url.lower().endswith(".docx")):
        return docx_doc(url, data)
    if "html" in ct or "xml" in ct or b"<html" in data[:2048].lower() or not ct:
        return html_doc(url, data, page.content_type)
    return Doc(url, "other", None, "", [])
